"""Valuations: median sold price per fragrance x size x condition, plus $/ml. Writes
docs/data.json (for the search page) and docs/fragrance_values.xlsx (for Excel)."""
import json
import statistics
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DOCS = Path(__file__).resolve().parent.parent / "docs"
NEW = {"new_sealed", "new_unsealed"}
USED = {"used", "partial", "tester", "unknown"}
SMALL = {"decant", "sample"}
SOURCES = ["reddit", "facebook", "mercari"]


def _cond_class(c):
    return "new" if c in NEW else "decant" if c in SMALL else "used"


def _drop_outliers(prices):
    if len(prices) < 5:
        return prices
    q = statistics.quantiles(prices, n=4)
    lo, hi = q[0] - 1.5 * (q[2] - q[0]), q[2] + 1.5 * (q[2] - q[0])
    return [p for p in prices if lo <= p <= hi]


def _stats(prices):
    p = _drop_outliers(sorted(prices))
    if not p:
        return None
    return {"median": round(statistics.median(p), 2), "low": round(min(p), 2), "high": round(max(p), 2), "n": len(p)}


def _date(ts):
    return datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d") if ts else None


def build(con, vcfg):
    now = time.time()
    window = now - vcfg["window_days"] * 86400
    recent_cut = now - 90 * 86400
    rows = [dict(r) for r in con.execute(
        "SELECT * FROM sales WHERE key IS NOT NULL AND price IS NOT NULL AND confidence>=? ORDER BY date_utc DESC",
        (vcfg["min_confidence"],),
    )]

    # display name per key = most common spelling Claude produced for it
    disp = {}
    for r in con.execute("SELECT key, brand, name, COUNT(*) n FROM sales WHERE key IS NOT NULL "
                         "GROUP BY key, brand, name ORDER BY n DESC"):
        disp.setdefault(r["key"], (r["brand"], r["name"]))
    for r in rows:
        r["brand"], r["name"] = disp.get(r["key"], (r["brand"], r["name"]))

    frags = {}
    for r in rows:
        f = frags.setdefault(r["key"], {"key": r["key"], "brand": r["brand"], "name": r["name"],
                                         "sold": [], "asks": []})
        (f["sold"] if r["status"] == "sold" else f["asks"] if r["status"] == "available" else []).append(r)

    out = []
    for f in frags.values():
        sold = [r for r in f["sold"] if not r["bundle"] and r["date_utc"] >= window]
        if not sold and not f["asks"]:
            continue
        groups = {}
        for r in sold:
            if r["size_ml"]:
                groups.setdefault((round(r["size_ml"]), _cond_class(r["condition"])), []).append(r)
        sizes = []
        for (size, cc), rs in sorted(groups.items()):
            s = _stats([r["price"] for r in rs])
            if not s:
                continue
            r90 = [r["price"] for r in rs if r["date_utc"] >= recent_cut]
            s["median_90d"] = round(statistics.median(r90), 2) if r90 else None
            s["avg_fill"] = round(statistics.mean([r["fill_pct"] for r in rs if r["fill_pct"] is not None]), 0) \
                if any(r["fill_pct"] is not None for r in rs) else None
            s["by_source"] = {}
            for src in SOURCES:
                st = _stats([r["price"] for r in rs if r["source"] == src])
                if st:
                    s["by_source"][src] = {"median": st["median"], "n": st["n"]}
            sizes.append({"size_ml": size, "cond": cc, **s})

        def ppm(rs):
            v = [r["price"] / (r["size_ml"] * r["fill_pct"] / 100) for r in rs
                 if r["size_ml"] and r["fill_pct"] and r["fill_pct"] > 0]
            s = _stats(v)
            return s["median"] if s else None

        asks = [r for r in f["asks"] if not r["bundle"] and r["date_utc"] >= recent_cut
                and _cond_class(r["condition"]) != "decant"]  # bottle asks only
        out.append({
            "key": f["key"], "brand": f["brand"], "name": f["name"],
            "n_sold": len(sold),
            "last_sold": _date(max((r["date_utc"] for r in sold), default=None)),
            "ppm_bottle": ppm([r for r in sold if _cond_class(r["condition"]) != "decant"]),
            "ppm_decant": ppm([r for r in sold if _cond_class(r["condition"]) == "decant"]),
            "sources": sorted({r["source"] for r in sold}),
            "n_by_source": {src: sum(1 for r in sold if r["source"] == src) for src in SOURCES},
            # compact rows so the search page can recompute values with platforms toggled off:
            # [days since epoch, price, size_ml, fill_pct, cond class, source index]
            "s": [[int(r["date_utc"] // 86400), r["price"], r["size_ml"], r["fill_pct"],
                   _cond_class(r["condition"])[0], SOURCES.index(r["source"])] for r in sold],
            "sizes": sizes,
            "ask_median": _stats([r["price"] for r in asks])["median"] if asks and _stats([r["price"] for r in asks]) else None,
            "n_asks": len(asks),
            "recent": [{
                "date": _date(r["date_utc"]), "price": r["price"], "size_ml": r["size_ml"], "fill": r["fill_pct"],
                "cond": r["condition"], "source": r["source"], "url": r["url"], "conc": r["concentration"],
            } for r in sold[:40]],
        })
    out.sort(key=lambda x: (-x["n_sold"], x["brand"] or "", x["name"] or ""))

    DOCS.mkdir(exist_ok=True)
    summary = {
        "updated": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "n_fragrances": len(out),
        "n_sales": sum(x["n_sold"] for x in out),
        "window_days": vcfg["window_days"],
        "sources": SOURCES,
        "default_off": vcfg.get("default_off") or [],
        "n_by_source": {src: sum(x["n_by_source"][src] for x in out) for src in SOURCES},
        "fragrances": out,
    }
    (DOCS / "data.json").write_text(json.dumps(summary, separators=(",", ":")))
    _excel(out, rows, summary)
    print(f"export: {len(out)} fragrances, {summary['n_sales']} sales -> docs/data.json, docs/fragrance_values.xlsx")


def _excel(out, rows, summary):
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B2F4A")

    def sheet(ws, headers, data, widths, money_cols=()):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in money_cols:
            for c in ws[get_column_letter(col)][1:]:
                c.number_format = '"$"#,##0.00'
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Values"
    data = []
    for f in out:
        for s in f["sizes"]:
            per = []
            for src in SOURCES:
                b = s["by_source"].get(src) or {}
                per += [b.get("median"), b.get("n") or 0]
            data.append([f["brand"], f["name"], s["size_ml"], s["cond"], s["avg_fill"], s["median"],
                         s["median_90d"], s["low"], s["high"], s["n"], *per, f["ppm_bottle"], f["last_sold"]])
    sheet(ws, ["Brand", "Fragrance", "Size (ml)", "Condition", "Avg fill %", "Median (all)", "Median 90d",
               "Low", "High", "# Sales", "Median Reddit", "# Reddit", "Median Facebook", "# Facebook",
               "Median Mercari", "# Mercari", "$/ml (bottles)", "Last sold"],
          data, [22, 30, 10, 11, 10, 13, 12, 10, 10, 9, 14, 9, 15, 11, 14, 10, 13, 12],
          money_cols=(6, 7, 8, 9, 11, 13, 15, 17))

    ws = wb.create_sheet("Sales")
    sheet(ws, ["Date", "Source", "Brand", "Fragrance", "Conc.", "Size (ml)", "Fill %", "Condition", "Price",
               "Status", "Bundle", "Confidence", "Evidence", "Link"],
          [[_date(r["date_utc"]), r["source"], r["brand"], r["name"], r["concentration"], r["size_ml"],
            r["fill_pct"], r["condition"], r["price"], r["status"], bool(r["bundle"]), r["confidence"],
            r["evidence"], r["url"]] for r in rows],
          [11, 9, 20, 28, 8, 9, 7, 12, 10, 10, 8, 10, 30, 40], money_cols=(9,))

    ws = wb.create_sheet("About")
    for line in [
        ["Fragrance values"], [f"Updated {summary['updated']}"],
        [f"{summary['n_fragrances']} fragrances from {summary['n_sales']} sold listings in the last {summary['window_days']} days"],
        ["Values = median SOLD price, outliers removed (1.5x IQR when 5+ sales). Sources: Reddit, Facebook groups, Mercari."],
        ["Per-platform medians are in the Values sheet; filter the Sales sheet's Source column to exclude a platform."],
        ["Condition: new = sealed/unsealed new; used = used/partial/tester; decant = decants & samples."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110
    wb.save(DOCS / "fragrance_values.xlsx")
